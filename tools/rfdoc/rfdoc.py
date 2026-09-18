#!/usr/bin/env python3
"""RF Doc (Site Build Form) generator.

Fills the RFDOC_template.xlsx sheet with site / sector RF data and writes
``<SITE_ID>_RFDOC.xlsx``.  Layout, styles and logo of the template are preserved.

Usage:
    rfdoc.py                    interactive prompts (defaults from --example values)
    rfdoc.py --gui              Tk form
    rfdoc.py --json site.json   fill from a JSON file (see --example)
    rfdoc.py --example          print the BM1183 example as JSON
    rfdoc.py --dump X_RFDOC.xlsx  read an existing RF Doc back to JSON
"""

from __future__ import annotations

import argparse
import copy
import datetime as dt
import json
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "RFDOC_template.xlsx"
SECTORS = ("A", "B", "C")
SECTOR_COLS = {"A": "B", "B": "D", "C": "F"}

# field -> (label, cell)
SITE_FIELDS = {
    "site_id": ("Site ID", "B7"),
    "site_address": ("Site address", "D7"),
    "east": ("East (WGS84)", "D9"),
    "north": ("North (WGS84)", "D10"),
    "rbs_type": ("RBS Type", "B13"),
    "bsc_rnc": ("BSC/RNC", "B43"),
    "comments": ("Comments", "B46"),
    "rp_engineer": ("RP Engineer", "B58"),
    "rp_manager": ("RP Manager", "E58"),
    "phone": ("Phone", "B61"),
    "date": ("Date (YYYY-MM-DD)", "E61"),
}

# field -> (label, row); column depends on sector
SECTOR_FIELDS = {
    "azimuth": ("Antenna Azimut", 17),
    "antenna_model": ("Antenna Model", 19),
    "n_antennas": ("Number of Antennas", 21),
    "hba": ("HBA", 23),
    "building_height": ("Height of the building", 25),
    "tower_height": ("Height of the Tower", 27),
    "mech_tilt": ("Mechanical Tilt", 29),
    "etilt_900": ("Electrical Tilt 900", 31),
    "etilt_1800": ("Electrical Tilt 1800", 33),
    "etilt_2100": ("Electrical Tilt 2100", 35),
    "trx": ("Number of TRX 900/1800", 37),
}

EXAMPLE = {
    "site_id": "BM1183",
    "site_address": "VILLAGE AGRICOLE CHAABET EL AMEUR BOUMERDES",
    "east": "3,714243 E",
    "north": "36,629465 N",
    "rbs_type": "2G/3G/4G Multi Band 900/1800/2100 (ITBBU V9200 R9264H B1/B3/B8)",
    "bsc_rnc": "BSCRBA/RNCRBA",
    "comments": "P12m",
    "rp_engineer": "Merzak  BRAHITI",
    "rp_manager": "Riad Mimoun",
    "phone": "541299477",
    "date": "2026-08-19",
    "sectors": {
        "A": {"azimuth": "60°"},
        "B": {"azimuth": "180°"},
        "C": {"azimuth": "270°"},
    },
}
_COMMON_SECTOR = {
    "antenna_model": "PB ODI-065R17M18JJJJ-GQ V",
    "n_antennas": "1",
    "hba": "8.5",
    "building_height": "10",
    "tower_height": "",
    "mech_tilt": "0°",
    "etilt_900": "2°",
    "etilt_1800": "2°",
    "etilt_2100": "2°",
    "trx": "2/0",
}
for _s in SECTORS:
    EXAMPLE["sectors"][_s] = {**EXAMPLE["sectors"][_s], **_COMMON_SECTOR}


def _num(v: str):
    """Store plain numbers as numbers so Excel can compute with them."""
    s = v.strip()
    if not s:
        return None
    try:
        return int(s)
    except ValueError:
        try:
            return float(s.replace(",", "."))
        except ValueError:
            return s


def _date(v: str):
    s = v.strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m-%d-%y"):
        try:
            return dt.datetime.strptime(s, fmt)
        except ValueError:
            pass
    return s


def fill(data: dict, out: Path | None = None) -> Path:
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb.active

    ws["B7"] = data.get("site_id", "")
    ws["D7"] = f"Site address:{data.get('site_address', '')}"
    ws["D9"] = f"East :  {data.get('east', '')}"
    ws["D10"] = f"North {data.get('north', '')}"
    ws["B13"] = data.get("rbs_type", "")
    ws["B43"] = data.get("bsc_rnc", "")
    ws["B46"] = data.get("comments", "")
    ws["B58"] = data.get("rp_engineer", "")
    ws["E58"] = data.get("rp_manager", "")
    ws["B61"] = _num(str(data.get("phone", "")))
    ws["E61"] = _date(str(data.get("date", "")))

    for sec in SECTORS:
        col = SECTOR_COLS[sec]
        values = data.get("sectors", {}).get(sec, {})
        for key, (_, row) in SECTOR_FIELDS.items():
            ws[f"{col}{row}"] = _num(str(values.get(key, "")))

    out = out or Path(f"{data.get('site_id') or 'SITE'}_RFDOC.xlsx")
    wb.save(out)
    return out


def dump(path: Path) -> dict:
    ws = openpyxl.load_workbook(path, data_only=True).active

    def s(cell):
        v = ws[cell].value
        if isinstance(v, dt.datetime):
            return v.strftime("%Y-%m-%d")
        return "" if v is None else str(v).strip()

    data = {
        "site_id": s("B7"),
        "site_address": s("D7").split(":", 1)[-1].strip(),
        "east": s("D9").split(":", 1)[-1].strip(),
        "north": s("D10").removeprefix("North").strip(),
        "rbs_type": s("B13"),
        "bsc_rnc": s("B43"),
        "comments": s("B46"),
        "rp_engineer": s("B58"),
        "rp_manager": s("E58"),
        "phone": s("B61"),
        "date": s("E61"),
        "sectors": {},
    }
    for sec in SECTORS:
        col = SECTOR_COLS[sec]
        data["sectors"][sec] = {
            k: s(f"{col}{row}") for k, (_, row) in SECTOR_FIELDS.items()
        }
    return data


def interactive() -> dict:
    data = copy.deepcopy(EXAMPLE)

    def ask(label, default):
        r = input(f"{label} [{default}]: ").strip()
        return r if r else default

    print("RF Doc - enter values (Enter keeps the default)\n")
    for key, (label, _) in SITE_FIELDS.items():
        data[key] = ask(label, data[key])
    for sec in SECTORS:
        print(f"\n--- Sector {sec} ---")
        for key, (label, _) in SECTOR_FIELDS.items():
            data["sectors"][sec][key] = ask(label, data["sectors"][sec][key])
    return data


def gui() -> None:
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk

    root = tk.Tk()
    root.title("RF Doc - Site Build Form")
    entries: dict[str, tk.Entry] = {}

    frm = ttk.Frame(root, padding=10)
    frm.grid(sticky="nsew")

    r = 0
    for key, (label, _) in SITE_FIELDS.items():
        ttk.Label(frm, text=label).grid(row=r, column=0, sticky="e", padx=4, pady=2)
        e = ttk.Entry(frm, width=70)
        e.insert(0, EXAMPLE[key])
        e.grid(row=r, column=1, columnspan=3, sticky="we", pady=2)
        entries[key] = e
        r += 1

    ttk.Separator(frm).grid(row=r, column=0, columnspan=4, sticky="we", pady=8)
    r += 1
    for i, sec in enumerate(SECTORS):
        ttk.Label(frm, text=f"Sector {sec}", font=("TkDefaultFont", 10, "bold")).grid(
            row=r, column=1 + i
        )
    r += 1
    for key, (label, _) in SECTOR_FIELDS.items():
        ttk.Label(frm, text=label).grid(row=r, column=0, sticky="e", padx=4, pady=2)
        for i, sec in enumerate(SECTORS):
            e = ttk.Entry(frm, width=24)
            e.insert(0, EXAMPLE["sectors"][sec][key])
            e.grid(row=r, column=1 + i, padx=2, pady=2)
            entries[f"{sec}.{key}"] = e
        r += 1

    def collect() -> dict:
        d = {k: entries[k].get() for k in SITE_FIELDS}
        d["sectors"] = {
            s: {k: entries[f"{s}.{k}"].get() for k in SECTOR_FIELDS} for s in SECTORS
        }
        return d

    def on_generate():
        d = collect()
        out = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            initialfile=f"{d['site_id'] or 'SITE'}_RFDOC.xlsx",
            filetypes=[("Excel", "*.xlsx")],
        )
        if not out:
            return
        fill(d, Path(out))
        messagebox.showinfo("RF Doc", f"Saved {out}")

    def on_load():
        p = filedialog.askopenfilename(
            filetypes=[("Excel", "*.xlsx"), ("JSON", "*.json")]
        )
        if not p:
            return
        d = json.loads(Path(p).read_text()) if p.endswith(".json") else dump(Path(p))
        for k in SITE_FIELDS:
            entries[k].delete(0, tk.END)
            entries[k].insert(0, d.get(k, ""))
        for s in SECTORS:
            for k in SECTOR_FIELDS:
                entries[f"{s}.{k}"].delete(0, tk.END)
                entries[f"{s}.{k}"].insert(
                    0, d.get("sectors", {}).get(s, {}).get(k, "")
                )

    btns = ttk.Frame(frm)
    btns.grid(row=r, column=0, columnspan=4, pady=10)
    ttk.Button(btns, text="Load existing RF Doc / JSON", command=on_load).pack(
        side="left", padx=5
    )
    ttk.Button(btns, text="Generate RF Doc", command=on_generate).pack(
        side="left", padx=5
    )
    root.mainloop()


def main(argv=None) -> int:
    p = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    p.add_argument("--json", type=Path, help="fill the form from a JSON file")
    p.add_argument("--gui", action="store_true", help="open the Tk form")
    p.add_argument("--example", action="store_true", help="print example JSON (BM1183)")
    p.add_argument(
        "--dump", type=Path, metavar="XLSX", help="read an RF Doc back to JSON"
    )
    p.add_argument(
        "-o", "--output", type=Path, help="output .xlsx (default <SITE_ID>_RFDOC.xlsx)"
    )
    a = p.parse_args(argv)

    if a.example:
        print(json.dumps(EXAMPLE, indent=2, ensure_ascii=False))
        return 0
    if a.dump:
        print(json.dumps(dump(a.dump), indent=2, ensure_ascii=False))
        return 0
    if a.gui:
        gui()
        return 0
    data = json.loads(a.json.read_text(encoding="utf-8")) if a.json else interactive()
    out = fill(data, a.output)
    print(f"Wrote {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
